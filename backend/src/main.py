from __future__ import annotations

import json
import hashlib
import secrets
import asyncio
from datetime import datetime, timedelta, timezone
from uuid import uuid4
from typing import Any

from fastapi import FastAPI, File, Form, Header, HTTPException, Request, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel, Field
from js import Object, Uint8Array, fetch
from pyodide.ffi import to_js
from workers import asgi

from services.excel_parser import ExcelImportError, ParsedQuestion, parse_question_excel


app = FastAPI(title="HHX English Math Classroom API", version="0.1.0")


app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:5173", "https://english-math-classroom.pages.dev"],
    allow_origin_regex=r"https://[a-z0-9-]+\.pages\.dev",
    allow_credentials=True,
    allow_methods=["GET", "POST", "OPTIONS"],
    allow_headers=["Authorization", "Content-Type"],
)


class LoginRequest(BaseModel):
    username: str = Field(min_length=1, max_length=80)
    password: str = Field(min_length=1, max_length=256)
    mode: str = Field(default="teacher", pattern="^(teacher|student)$")


class StudentRoundStartRequest(BaseModel):
    topic_id: str | None = None
    access_code: str | None = None


class StudentRoundAnswerItem(BaseModel):
    question_id: str = Field(min_length=1)
    selected_option: str | None = None
    text_answer: str | None = None


class StudentRoundAnswerRequest(BaseModel):
    round_id: str = Field(min_length=1)
    question_id: str = Field(min_length=1)
    selected_option: str | None = None
    text_answer: str | None = None
    confirm: bool = False


class StudentRoundFinishRequest(BaseModel):
    round_id: str = Field(min_length=1)
    answers: list[StudentRoundAnswerItem] = Field(default_factory=list)


class ClassroomAnswerRequest(BaseModel):
    question_id: str
    selected_option: str = Field(pattern="^[ABCD]$")


class WordRatingRequest(BaseModel):
    word_id: str
    rating: str = Field(pattern="^(forgot|fuzzy|clear)$")


def utc_now() -> datetime:
    return datetime.now(timezone.utc)


def request_env(request: Request) -> Any:
    return request.scope.get("env")


def env_value(request: Request, name: str, default: str = "") -> str:
    runtime_env = request_env(request)
    value = getattr(runtime_env, name, None) if runtime_env is not None else None
    return str(value or default)


async def create_student_session(request: Request, student: dict[str, Any]) -> tuple[str, datetime]:
    token = secrets.token_urlsafe(32)
    expires_at = utc_now() + timedelta(hours=24)
    status, result = await supabase_request(
        request,
        "POST",
        "/rest/v1/student_sessions",
        service_role=True,
        prefer_representation=True,
        body={
            "student_id": student["student_id"],
            "token_hash": hashlib.sha256(token.encode("utf-8")).hexdigest(),
            "auth_version": student["auth_version"],
            "expires_at": expires_at.isoformat(),
        },
    )
    if status >= 400 or not result:
        raise HTTPException(status_code=502, detail="無法建立學生登入狀態")
    return token, expires_at


async def verify_student_token(request: Request, authorization: str | None) -> dict[str, Any]:
    token = bearer_token(authorization)
    token_hash = hashlib.sha256(token.encode("utf-8")).hexdigest()
    status, sessions = await supabase_request(
        request,
        "GET",
        "/rest/v1/student_sessions",
        service_role=True,
        params={
            "token_hash": f"eq.{token_hash}",
            "revoked_at": "is.null",
            "expires_at": f"gt.{utc_now().isoformat()}",
            "select": "student_id,auth_version,students(id,name,student_number,school_id,grade_id,class_id,auth_version,is_active)",
        },
    )
    if status >= 400 or not sessions:
        raise HTTPException(status_code=401, detail="學生登入狀態已失效，請重新登入")
    session = sessions[0]
    student = session.get("students")
    if not student or not student.get("is_active") or student.get("auth_version") != session.get("auth_version"):
        raise HTTPException(status_code=401, detail="學生登入狀態已失效，請重新登入")
    return student


def supabase_headers(
    request: Request,
    access_token: str | None = None,
    *,
    service_role: bool = False,
    prefer_representation: bool = False,
) -> dict[str, str]:
    if service_role:
        api_key = env_value(request, "SUPABASE_SECRET_KEY") or env_value(request, "SUPABASE_SERVICE_ROLE_KEY")
    else:
        api_key = env_value(request, "SUPABASE_ANON_KEY")
    if not api_key:
        raise HTTPException(status_code=503, detail="後端尚未設定 Supabase API Key")
    headers = {"apikey": api_key, "Content-Type": "application/json"}
    if service_role:
        headers["Authorization"] = f"Bearer {api_key}"
    if access_token:
        headers["Authorization"] = f"Bearer {access_token}"
    if prefer_representation:
        headers["Prefer"] = "return=representation"
    return headers


def supabase_url(request: Request) -> str:
    value = env_value(request, "SUPABASE_URL").rstrip("/")
    if not value:
        raise HTTPException(status_code=503, detail="後端尚未設定 Supabase URL")
    return value


async def supabase_request(
    request: Request,
    method: str,
    path: str,
    *,
    access_token: str | None = None,
    body: dict[str, Any] | None = None,
    params: dict[str, str] | None = None,
    service_role: bool = False,
    prefer_representation: bool = False,
) -> tuple[int, Any]:
    last_error: Exception | None = None
    for attempt in range(2):
        try:
            url = f"{supabase_url(request)}{path}"
            if params:
                from urllib.parse import urlencode
                url = f"{url}?{urlencode(params)}"
            database_request = path.startswith("/rest/v1/")
            use_service_role = service_role or database_request
            request_headers = supabase_headers(
                request,
                None if use_service_role else access_token,
                service_role=use_service_role,
                prefer_representation=prefer_representation,
            )
            request_init: dict[str, Any] = {"method": method, "headers": request_headers}
            if body is not None:
                request_init["body"] = json.dumps(body)
            response = await fetch(url, to_js(request_init, dict_converter=Object.fromEntries))
            response_text = str(await response.text())
            try:
                payload = json.loads(response_text)
            except json.JSONDecodeError:
                payload = {"message": response_text[:300]}
            return int(response.status), payload
        except Exception as exc:
            last_error = exc
            if attempt == 1:
                return 599, {"message": str(exc)[:300]}
            await asyncio.sleep(0.05)
    return 599, {"message": str(last_error)[:300] if last_error else "Supabase request failed"}


def bearer_token(authorization: str | None) -> str:
    if not authorization or not authorization.lower().startswith("bearer "):
        raise HTTPException(status_code=401, detail="請先登入")
    token = authorization[7:].strip()
    if not token:
        raise HTTPException(status_code=401, detail="登入狀態無效")
    return token


async def current_teacher_profile(request: Request, token: str) -> dict[str, Any]:
    user_status, user = await supabase_request(request, "GET", "/auth/v1/user", access_token=token)
    if user_status >= 400 or not user.get("id"):
        raise HTTPException(status_code=401, detail="登入狀態已失效，請重新登入")
    profile_status, profiles = await supabase_request(
        request,
        "GET",
        "/rest/v1/teacher_profiles",
        service_role=True,
        params={"id": f"eq.{user['id']}", "select": "id,display_name,school_id,is_active"},
    )
    if profile_status >= 400 or not profiles:
        raise HTTPException(status_code=403, detail="此帳號尚未建立老師資料")
    profile = profiles[0]
    if not profile.get("is_active"):
        raise HTTPException(status_code=403, detail="此老師帳號已停用")
    return {"user_id": user["id"], **profile}


def import_error_response(error: ExcelImportError) -> HTTPException:
    return HTTPException(status_code=422, detail={"message": str(error), "errors": error.errors})


async def _storage_upload(request: Request, path: str, data_url: str) -> tuple[str, int]:
    import base64
    from urllib.parse import quote

    try:
        header, encoded = data_url.split(",", 1)
        mime_type = header.split(";")[0].removeprefix("data:")
        raw = base64.b64decode(encoded, validate=True)
    except Exception as exc:
        raise HTTPException(status_code=422, detail="题目图片格式无效") from exc
    if mime_type not in {"image/jpeg", "image/png", "image/webp"} or not raw or len(raw) > 1024 * 1024:
        raise HTTPException(status_code=422, detail="题目图片必须是 1MB 以内的 JPG、PNG 或 WebP")
    key = env_value(request, "SUPABASE_SECRET_KEY") or env_value(request, "SUPABASE_SERVICE_ROLE_KEY")
    if not key:
        raise HTTPException(status_code=503, detail="后端尚未设置 Supabase Service Role Key")
    js_bytes = Uint8Array.new(len(raw))
    for index, value in enumerate(raw):
        js_bytes[index] = value
    init = to_js({"method": "POST", "headers": {"apikey": key, "Authorization": f"Bearer {key}", "Content-Type": mime_type, "x-upsert": "true"}, "body": js_bytes}, dict_converter=Object.fromEntries)
    response = await fetch(f"{supabase_url(request)}/storage/v1/object/question-images-v2/{quote(path, safe='/')}", init)
    if int(response.status) >= 400:
        raise HTTPException(status_code=502, detail="题目图片上传失败")
    return path, len(raw)


async def _signed_image_url(request: Request, path: str) -> str:
    from urllib.parse import quote
    status, result = await supabase_request(request, "POST", f"/storage/v1/object/sign/question-images-v2/{quote(path, safe='/')}", service_role=True, body={"expiresIn": 3600})
    signed = result.get("signedURL") or result.get("signedUrl") if isinstance(result, dict) else None
    if status >= 400 or not signed:
        return ""
    return signed if signed.startswith("http") else f"{supabase_url(request)}{signed}"


@app.get("/api/teacher/topics")
async def teacher_topics(request: Request, authorization: str | None = Header(default=None)) -> list[dict[str, Any]]:
    token = bearer_token(authorization)
    teacher = await current_teacher_profile(request, token)
    status, topics = await supabase_request(request, "GET", "/rest/v1/topics", service_role=True, params={"school_id": f"eq.{teacher['school_id']}", "is_active": "eq.true", "select": "id,code,name,sort_order,math_batches(id,status,created_at,published_at,questions(count))", "order": "sort_order.asc"})
    if status >= 400:
        raise HTTPException(status_code=502, detail="无法读取主题目录")
    rows = []
    for topic in topics:
        batches = topic.pop("math_batches", [])
        published = next((batch for batch in batches if batch["status"] == "published"), None)
        rows.append({**topic, "batch_id": published["id"] if published else None, "batch_status": published["status"] if published else "empty", "question_count": (published.get("questions") or [{"count": 0}])[0].get("count", 0) if published else 0})
    return rows


@app.get("/api/teacher/topics/{topic_id}")
async def teacher_topic_detail(topic_id: str, request: Request, authorization: str | None = Header(default=None)) -> dict[str, Any]:
    token = bearer_token(authorization)
    teacher = await current_teacher_profile(request, token)
    status, topics = await supabase_request(request, "GET", "/rest/v1/topics", service_role=True, params={"id": f"eq.{topic_id}", "school_id": f"eq.{teacher['school_id']}", "select": "id,code,name"})
    if status >= 400 or not topics:
        raise HTTPException(status_code=404, detail="找不到主题")
    _, batches = await supabase_request(request, "GET", "/rest/v1/math_batches", service_role=True, params={"topic_id": f"eq.{topic_id}", "status": "eq.published", "select": "id,created_at,published_at,questions(id,sort_order,question_type,image_bucket,image_path,correct_option,source_reference,source_year,source_question_number,source_paper)", "order": "created_at.desc", "limit": "1"})
    batch = batches[0] if batches else None
    questions = []
    if batch:
        for q in sorted(batch.get("questions", []), key=lambda item: item["sort_order"]):
            questions.append({**q, "question_image_url": await _signed_question_image(request, q.get("image_bucket", "question-images-v2"), q["image_path"])})
    return {**topics[0], "batch": batch, "questions": questions}


@app.get("/api/teacher/questions")
async def search_teacher_questions(request: Request, year: str | None = None, authorization: str | None = Header(default=None)) -> list[dict[str, Any]]:
    token = bearer_token(authorization)
    teacher = await current_teacher_profile(request, token)
    params = {"source_year": f"eq.{year.strip()}", "select": "id,batch_id,sort_order,source_reference,source_year,source_question_number,source_paper,math_batches!inner(topic_id,topics!inner(code,name,school_id))", "math_batches.topics.school_id": f"eq.{teacher['school_id']}", "order": "source_year.asc,source_question_number.asc"} if year and year.strip() else {"select": "id,batch_id,sort_order,source_reference,source_year,source_question_number,source_paper,math_batches!inner(topic_id,topics!inner(code,name,school_id))", "math_batches.topics.school_id": f"eq.{teacher['school_id']}", "order": "source_year.asc,source_question_number.asc"}
    status, rows = await supabase_request(request, "GET", "/rest/v1/questions", service_role=True, params=params)
    if status >= 400:
        raise HTTPException(status_code=502, detail="无法检索题目来源")
    return rows


@app.post("/api/teacher/topics/preview")
async def preview_topic_import(request: Request, topic_id: str = Form(...), file: UploadFile = File(...), authorization: str | None = Header(default=None)) -> dict[str, Any]:
    token = bearer_token(authorization)
    teacher = await current_teacher_profile(request, token)
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=422, detail="只支持 .xlsx 格式的 Excel 文件")
    topic_status, topics = await supabase_request(request, "GET", "/rest/v1/topics", service_role=True, params={"id": f"eq.{topic_id}", "school_id": f"eq.{teacher['school_id']}", "select": "id"})
    if topic_status >= 400 or not topics:
        raise HTTPException(status_code=404, detail="找不到主题")
    try:
        questions = parse_question_excel(await file.read())
    except ExcelImportError as error:
        raise import_error_response(error) from error
    return {"filename": file.filename, "topic_id": topic_id, "question_count": len(questions), "questions": [{"row_number": q.row_number, "question_text": q.question_text, "question_image_url": q.question_image_url, "question_type": q.question_type, "correct_answer": q.correct_answer, "source_reference": q.source_reference, "source_year": q.source_year, "source_question_number": q.source_question_number, "source_paper": q.source_paper} for q in questions]}


@app.post("/api/teacher/topics/import")
async def import_topic(request: Request, topic_id: str = Form(...), file: UploadFile = File(...), authorization: str | None = Header(default=None)) -> dict[str, Any]:
    token = bearer_token(authorization)
    teacher = await current_teacher_profile(request, token)
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=422, detail="只支持 .xlsx 格式的 Excel 文件")
    topic_status, topics = await supabase_request(request, "GET", "/rest/v1/topics", service_role=True, params={"id": f"eq.{topic_id}", "school_id": f"eq.{teacher['school_id']}", "select": "id,code"})
    if topic_status >= 400 or not topics:
        raise HTTPException(status_code=404, detail="找不到主题")
    try:
        questions = parse_question_excel(await file.read())
    except ExcelImportError as error:
        raise import_error_response(error) from error
    job_id, batch_id = str(uuid4()), str(uuid4())
    status, _ = await supabase_request(request, "POST", "/rest/v1/import_jobs", service_role=True, body={"id": job_id, "created_by": teacher["user_id"], "kind": "math", "topic_id": topic_id, "filename": file.filename, "status": "preview"})
    if status >= 400:
        raise HTTPException(status_code=502, detail="无法建立导入记录")
    status, _ = await supabase_request(request, "POST", "/rest/v1/math_batches", service_role=True, body={"id": batch_id, "topic_id": topic_id, "created_by": teacher["user_id"], "import_job_id": job_id, "status": "draft"})
    if status >= 400:
        raise HTTPException(status_code=502, detail="无法建立题库批次")
    question_rows = []
    for order, q in enumerate(questions, start=1):
        if not q.question_image_url:
            await supabase_request(request, "PATCH", "/rest/v1/import_jobs", service_role=True, params={"id": f"eq.{job_id}"}, body={"status": "failed", "errors": [{"row": q.row_number, "message": "题目截图不能为空"}]})
            raise HTTPException(status_code=422, detail=f"第 {q.row_number} 行缺少题目截图")
        image_path = f"{topic_id}/{batch_id}/{order}-{secrets.token_hex(6)}.jpg"
        _, image_bytes = await _storage_upload(request, image_path, q.question_image_url)
        question_rows.append({"id": str(uuid4()), "batch_id": batch_id, "sort_order": order, "question_type": q.question_type, "image_path": image_path, "image_bytes": image_bytes, "correct_option": q.correct_answer, "source_reference": q.source_reference, "source_year": q.source_year, "source_question_number": q.source_question_number, "source_paper": q.source_paper})
    status, result = await supabase_request(request, "POST", "/rest/v1/questions", service_role=True, prefer_representation=True, body=question_rows)
    if status >= 400:
        await supabase_request(request, "PATCH", "/rest/v1/import_jobs", service_role=True, params={"id": f"eq.{job_id}"}, body={"status": "failed", "errors": [{"message": "写入题目失败"}]})
        raise HTTPException(status_code=502, detail="题目写入失败")
    status, _ = await supabase_request(request, "POST", "/rest/v1/rpc/app_publish_math_batch", service_role=True, body={"p_batch": batch_id, "p_teacher": teacher["user_id"]})
    if status >= 400:
        await supabase_request(request, "PATCH", "/rest/v1/import_jobs", service_role=True, params={"id": f"eq.{job_id}"}, body={"status": "failed", "errors": [{"message": "发布批次失败"}]})
        raise HTTPException(status_code=422, detail="题目已上传但发布失败，请联系管理员检查导入记录")
    return {"topic_id": topic_id, "batch_id": batch_id, "name": topics[0]["code"], "question_count": len(result)}


async def _classroom_payload(request: Request, session: dict[str, Any], *, include_answer: bool = False) -> dict[str, Any]:
    session = await _advance_classroom_if_expired(request, session)
    if not session.get("batch_id"):
        status, rows = await supabase_request(request, "GET", "/rest/v1/classroom_sessions", service_role=True, params={"id": f"eq.{session['id']}", "select": "id,batch_id,class_id,status,current_question_id,question_status,duration_seconds,question_opened_at,deadline_at"})
        if status < 400 and rows:
            session = rows[0]
    question = None
    if session.get("current_question_id"):
        status, rows = await supabase_request(request, "GET", "/rest/v1/questions", service_role=True, params={"id": f"eq.{session['current_question_id']}", "batch_id": f"eq.{session['batch_id']}", "select": "id,sort_order,question_type,image_bucket,image_path"})
        if rows:
            row = rows[0]
            question = {"id": row["id"], "sort_order": row["sort_order"], "question_type": row["question_type"], "image_url": await _signed_question_image(request, row["image_bucket"], row["image_path"])}
    participant_status, participants = await supabase_request(request, "GET", "/rest/v1/classroom_participants", service_role=True, params={"session_id": f"eq.{session['id']}", "select": "student_id"})
    answer_count = 0
    if question:
        _, answers = await supabase_request(request, "GET", "/rest/v1/classroom_answers", service_role=True, params={"session_id": f"eq.{session['id']}", "question_id": f"eq.{question['id']}", "select": "selected_option,is_correct"})
        answer_count = len(answers)
        if include_answer:
            question["distribution"] = {key: sum(1 for answer in answers if answer["selected_option"] == key) for key in "ABCD"}
    return {**session, "question": question, "participant_count": len(participants) if participant_status < 400 else 0, "submitted_count": answer_count}


async def _advance_classroom_if_expired(request: Request, session: dict[str, Any]) -> dict[str, Any]:
    if session.get("status") != "active" or not session.get("deadline_at"):
        return session
    try:
        deadline = datetime.fromisoformat(session["deadline_at"].replace("Z", "+00:00"))
    except (TypeError, ValueError):
        return session
    if deadline > utc_now():
        return session
    status, links = await supabase_request(
        request, "GET", "/rest/v1/classroom_session_questions", service_role=True,
        params={"session_id": f"eq.{session['id']}", "select": "question_id,sort_order", "order": "sort_order.asc"},
    )
    if status >= 400:
        return session
    index = next((i for i, item in enumerate(links) if item["question_id"] == session.get("current_question_id")), -1)
    now = utc_now()
    duration = int(session.get("duration_seconds") or 60)
    if index < 0 or index + 1 >= len(links):
        body = {"status": "closed", "question_status": "locked", "closed_at": now.isoformat(), "question_opened_at": None, "deadline_at": None}
    else:
        body = {"status": "active", "current_question_id": links[index + 1]["question_id"], "question_status": "open", "question_opened_at": now.isoformat(), "deadline_at": (now + timedelta(seconds=duration)).isoformat()}
    update_status, updated = await supabase_request(
        request, "PATCH", "/rest/v1/classroom_sessions", service_role=True, prefer_representation=True,
        params={"id": f"eq.{session['id']}", "current_question_id": f"eq.{session.get('current_question_id')}"}, body=body,
    )
    return updated[0] if update_status < 400 and updated else session


@app.get("/api/classroom/{access_token}")
async def get_classroom(access_token: str, request: Request, authorization: str | None = Header(default=None)) -> dict[str, Any]:
    await verify_student_token(request, authorization)
    status, rows = await supabase_request(request, "GET", "/rest/v1/classroom_sessions", service_role=True, params={"access_token": f"eq.{access_token}", "select": "id,batch_id,class_id,access_token,status,current_question_id,question_status,duration_seconds,question_opened_at,deadline_at,expires_at,closed_at", "limit": "1"})
    if status >= 400 or not rows or rows[0]["status"] in {"closed", "invalidated"}:
        raise HTTPException(status_code=404, detail="课堂已结束或链接失效")
    return await _classroom_payload(request, rows[0])


@app.post("/api/classroom/{access_token}/join")
async def join_classroom(access_token: str, request: Request, authorization: str | None = Header(default=None)) -> dict[str, Any]:
    student = await verify_student_token(request, authorization)
    status, rows = await supabase_request(request, "GET", "/rest/v1/classroom_sessions", service_role=True, params={"access_token": f"eq.{access_token}", "select": "id,batch_id,status,expires_at"})
    if status >= 400 or not rows or rows[0]["status"] in {"closed", "invalidated"}:
        raise HTTPException(status_code=404, detail="课堂已结束或链接失效")
    session = rows[0]
    existing_status, existing = await supabase_request(request, "GET", "/rest/v1/classroom_participants", service_role=True, params={"session_id": f"eq.{session['id']}", "student_id": f"eq.{student['id']}", "select": "session_id"})
    if existing_status >= 400:
        raise HTTPException(status_code=502, detail="无法读取课堂参与状态")
    if not existing:
        join_status, _ = await supabase_request(request, "POST", "/rest/v1/classroom_participants", service_role=True, body={"session_id": session["id"], "student_id": student["id"]})
        if join_status >= 400 and join_status != 409:
            raise HTTPException(status_code=409, detail="无法加入课堂")
    return {"joined": True}


@app.post("/api/classroom/{access_token}/answer")
async def answer_classroom(access_token: str, payload: ClassroomAnswerRequest, request: Request, authorization: str | None = Header(default=None)) -> dict[str, Any]:
    student = await verify_student_token(request, authorization)
    status, rows = await supabase_request(request, "GET", "/rest/v1/classroom_sessions", service_role=True, params={"access_token": f"eq.{access_token}", "select": "id,batch_id,status,current_question_id,question_status"})
    if status >= 400 or not rows or rows[0]["status"] != "active" or rows[0]["question_status"] != "open" or rows[0]["current_question_id"] != payload.question_id:
        raise HTTPException(status_code=409, detail="课堂题目已关闭")
    answer_params = {"session_id": f"eq.{rows[0]['id']}", "student_id": f"eq.{student['id']}", "question_id": f"eq.{payload.question_id}"}
    existing_status, existing = await supabase_request(request, "GET", "/rest/v1/classroom_answers", service_role=True, params={**answer_params, "select": "session_id"})
    if existing_status >= 400:
        raise HTTPException(status_code=502, detail="无法读取答题状态")
    if existing:
        answer_status, result = await supabase_request(request, "PATCH", "/rest/v1/classroom_answers", service_role=True, prefer_representation=True, params=answer_params, body={"selected_option": payload.selected_option})
    else:
        answer_status, result = await supabase_request(request, "POST", "/rest/v1/classroom_answers", service_role=True, prefer_representation=True, body={"session_id": rows[0]["id"], "student_id": student["id"], "question_id": payload.question_id, "selected_option": payload.selected_option, "is_correct": False})
    if answer_status >= 400:
        raise HTTPException(status_code=409, detail="答案提交失败")
    return {"submitted": True}


@app.post("/api/teacher/classrooms/{session_id}/start")
async def start_new_classroom(session_id: str, request: Request, authorization: str | None = Header(default=None)) -> dict[str, Any]:
    token = bearer_token(authorization); teacher = await current_teacher_profile(request, token)
    status, rows = await supabase_request(request, "GET", "/rest/v1/classroom_sessions", service_role=True, params={"id": f"eq.{session_id}", "created_by": f"eq.{teacher['user_id']}", "select": "id,batch_id,status,duration_seconds"})
    if status >= 400 or not rows: raise HTTPException(status_code=404, detail="找不到课堂")
    _, questions = await supabase_request(request, "GET", "/rest/v1/classroom_session_questions", service_role=True, params={"session_id": f"eq.{session_id}", "select": "question_id,sort_order", "order": "sort_order.asc"})
    if not questions: raise HTTPException(status_code=409, detail="课堂没有题目")
    now = utc_now(); duration = int(rows[0].get("duration_seconds") or 60); update_status, result = await supabase_request(request, "PATCH", "/rest/v1/classroom_sessions", service_role=True, prefer_representation=True, params={"id": f"eq.{session_id}"}, body={"status": "active", "current_question_id": questions[0]["question_id"], "question_status": "open", "question_opened_at": now.isoformat(), "deadline_at": (now + timedelta(seconds=duration)).isoformat()})
    if update_status >= 400: raise HTTPException(status_code=502, detail="课堂启动失败")
    return result[0]


@app.post("/api/teacher/classrooms/{session_id}/next")
async def next_new_classroom(session_id: str, request: Request, authorization: str | None = Header(default=None)) -> dict[str, Any]:
    token = bearer_token(authorization); teacher = await current_teacher_profile(request, token)
    status, rows = await supabase_request(request, "GET", "/rest/v1/classroom_sessions", service_role=True, params={"id": f"eq.{session_id}", "created_by": f"eq.{teacher['user_id']}", "select": "id,status,current_question_id,duration_seconds"})
    if status >= 400 or not rows: raise HTTPException(status_code=404, detail="找不到课堂")
    session = rows[0]
    _, questions = await supabase_request(request, "GET", "/rest/v1/classroom_session_questions", service_role=True, params={"session_id": f"eq.{session_id}", "select": "question_id,sort_order", "order": "sort_order.asc"})
    index = next((i for i, q in enumerate(questions) if q["question_id"] == session.get("current_question_id")), -1)
    now = utc_now(); duration = int(session.get("duration_seconds") or 60)
    if index < 0 or index + 1 >= len(questions):
        body = {"status": "closed", "question_status": "locked", "closed_at": now.isoformat(), "question_opened_at": None, "deadline_at": None}
    else:
        body = {"status": "active", "current_question_id": questions[index + 1]["question_id"], "question_status": "open", "question_opened_at": now.isoformat(), "deadline_at": (now + timedelta(seconds=duration)).isoformat()}
    updated_status, result = await supabase_request(request, "PATCH", "/rest/v1/classroom_sessions", service_role=True, prefer_representation=True, params={"id": f"eq.{session_id}"}, body=body)
    if updated_status >= 400: raise HTTPException(status_code=502, detail="切换课堂题目失败")
    return result[0]


@app.get("/api/teacher/classrooms/{session_id}/stats")
async def new_classroom_stats(session_id: str, request: Request, authorization: str | None = Header(default=None)) -> dict[str, Any]:
    token = bearer_token(authorization); teacher = await current_teacher_profile(request, token)
    status, rows = await supabase_request(request, "GET", "/rest/v1/classroom_sessions", service_role=True, params={"id": f"eq.{session_id}", "created_by": f"eq.{teacher['user_id']}", "select": "id,batch_id,class_id,status,current_question_id,question_status,duration_seconds,question_opened_at,deadline_at"})
    if status >= 400 or not rows: raise HTTPException(status_code=404, detail="找不到课堂")
    return await _classroom_payload(request, rows[0], include_answer=True)


@app.get("/api/teacher/reports")
async def teacher_report_index(request: Request, authorization: str | None = Header(default=None)) -> list[dict[str, Any]]:
    token = bearer_token(authorization); teacher = await current_teacher_profile(request, token)
    status, classrooms = await supabase_request(request, "GET", "/rest/v1/classroom_sessions", service_role=True, params={"created_by": f"eq.{teacher['user_id']}", "select": "id,batch_id,class_id,status,created_at,closed_at,classes(name,code),math_batches(topics(code,name))", "order": "created_at.desc"})
    practice_status, practices = await supabase_request(request, "GET", "/rest/v1/practices", service_role=True, params={"created_by": f"eq.{teacher['user_id']}", "select": "id,batch_id,access_code,status,created_at,math_batches(topics(code,name))", "order": "created_at.desc"})
    result = []
    for item in classrooms if status < 400 else []:
        topic = (item.get("math_batches") or {}).get("topics") or {}
        cls = item.get("classes") or {}
        result.append({"id": item["id"], "kind": "classroom", "status": item["status"], "title": f"{topic.get('code','')} {topic.get('name','')}", "subtitle": cls.get("name") or cls.get("code", ""), "created_at": item.get("created_at"), "closed_at": item.get("closed_at")})
    for item in practices if practice_status < 400 else []:
        topic = (item.get("math_batches") or {}).get("topics") or {}
        result.append({"id": item["id"], "kind": "practice", "status": item["status"], "title": f"{topic.get('code','')} {topic.get('name','')}", "subtitle": f"练习码 {item.get('access_code','')}", "created_at": item.get("created_at")})
    return sorted(result, key=lambda row: row.get("created_at") or "", reverse=True)


@app.get("/api/teacher/reports/practice-summary")
async def teacher_practice_summary(
    request: Request,
    class_id: str,
    batch_id: str,
    authorization: str | None = Header(default=None),
) -> dict[str, Any]:
    token = bearer_token(authorization)
    teacher = await current_teacher_profile(request, token)
    class_status, classes = await supabase_request(
        request, "GET", "/rest/v1/classes", service_role=True,
        params={"id": f"eq.{class_id}", "school_id": f"eq.{teacher['school_id']}", "is_active": "eq.true", "select": "id,code,name,grade_id,grades(code,name)"},
    )
    if class_status >= 400 or not classes:
        raise HTTPException(status_code=404, detail="找不到可用班级")
    batch_status, batches = await supabase_request(
        request, "GET", "/rest/v1/math_batches", service_role=True,
        params={"id": f"eq.{batch_id}", "status": "eq.published", "select": "id,topic_id,topics!inner(code,name,school_id)"},
    )
    if batch_status >= 400 or not batches or batches[0].get("topics", {}).get("school_id") != teacher["school_id"]:
        raise HTTPException(status_code=404, detail="找不到当前题库")
    students_status, students = await supabase_request(
        request, "GET", "/rest/v1/students", service_role=True,
        params={"class_id": f"eq.{class_id}", "is_active": "eq.true", "select": "id,name,student_number", "order": "student_number.asc"},
    )
    if students_status >= 400:
        raise HTTPException(status_code=502, detail="无法读取班级学生")
    student_ids = [student["id"] for student in students]
    rounds: list[dict[str, Any]] = []
    if student_ids:
        round_status, rounds = await supabase_request(
            request, "GET", "/rest/v1/practice_rounds", service_role=True,
            params={"batch_id": f"eq.{batch_id}", "student_id": f"in.({','.join(student_ids)})", "status": "eq.completed", "select": "id,student_id,completed_at", "order": "completed_at.desc"},
        )
        if round_status >= 400:
            raise HTTPException(status_code=502, detail="无法读取班级答题记录")
    latest_round_ids: list[str] = []
    latest_round_by_student: dict[str, dict[str, Any]] = {}
    round_by_id: dict[str, dict[str, Any]] = {}
    for round_row in rounds:
        if round_row.get("id"):
            round_by_id[round_row["id"]] = round_row
        student_id = round_row.get("student_id")
        if student_id and student_id not in latest_round_by_student:
            latest_round_by_student[student_id] = round_row
            latest_round_ids.append(round_row["id"])
    answers: list[dict[str, Any]] = []
    if latest_round_ids:
        answer_status, answers = await supabase_request(
            request, "GET", "/rest/v1/practice_answers", service_role=True,
            params={"round_id": f"in.({','.join(latest_round_ids)})", "state": "in.(confirmed,unanswered)", "select": "question_id,selected_option,text_answer,is_correct,round_id,updated_at"},
        )
        if answer_status >= 400:
            raise HTTPException(status_code=502, detail="无法读取题目答案")
    question_status, questions = await supabase_request(
        request, "GET", "/rest/v1/questions", service_role=True,
        params={"batch_id": f"eq.{batch_id}", "select": "id,sort_order,question_type,correct_option", "order": "sort_order.asc"},
    )
    if question_status >= 400:
        raise HTTPException(status_code=502, detail="无法读取题库题目")
    # 课堂和自主练习都落在不同的答案表；按学生和题目取最新一次，避免重复计数。
    merged: dict[tuple[str, str], dict[str, Any]] = {}
    for answer in answers:
        round_row = round_by_id.get(answer.get("round_id"))
        if not round_row:
            continue
        student_id = round_row.get("student_id")
        key = (student_id, answer.get("question_id"))
        merged[key] = {**answer, "student_id": student_id, "answered_at": answer.get("updated_at") or round_row.get("completed_at") or ""}
    if student_ids:
        classroom_status, sessions = await supabase_request(
            request, "GET", "/rest/v1/classroom_sessions", service_role=True,
            params={"batch_id": f"eq.{batch_id}", "class_id": f"eq.{class_id}", "created_by": f"eq.{teacher['user_id']}", "select": "id"},
        )
        if classroom_status >= 400:
            raise HTTPException(status_code=502, detail="无法读取课堂记录")
        session_ids = [row["id"] for row in sessions]
        if session_ids:
            classroom_answer_status, classroom_answers = await supabase_request(
                request, "GET", "/rest/v1/classroom_answers", service_role=True,
                params={"session_id": f"in.({','.join(session_ids)})", "select": "student_id,question_id,selected_option,is_correct,answered_at", "order": "answered_at.asc"},
            )
            if classroom_answer_status >= 400:
                raise HTTPException(status_code=502, detail="无法读取课堂答案")
            for answer in classroom_answers:
                key = (answer.get("student_id"), answer.get("question_id"))
                previous = merged.get(key)
                if not previous or (answer.get("answered_at") or "") >= (previous.get("answered_at") or ""):
                    merged[key] = answer
            participant_status, participants = await supabase_request(
                request, "GET", "/rest/v1/classroom_participants", service_role=True,
                params={"session_id": f"in.({','.join(session_ids)})", "select": "student_id"},
            )
            if participant_status < 400:
                merged_participant_ids = {row.get("student_id") for row in participants if row.get("student_id")}
            else:
                merged_participant_ids = set()
        else:
            merged_participant_ids = set()
    else:
        merged_participant_ids = set()
    participant_ids = {student_id for student_id, _ in merged} | merged_participant_ids
    profile_map: dict[str, dict[str, Any]] = {}
    if participant_ids:
        profile_status, profiles = await supabase_request(
            request, "GET", "/rest/v1/students", service_role=True,
            params={"id": f"in.({','.join(participant_ids)})", "select": "id,name,student_number,classes(name,code)"},
        )
        if profile_status < 400:
            profile_map = {row["id"]: row for row in profiles}
    reports = []
    for question in questions:
        current = [answer for answer in merged.values() if answer.get("question_id") == question["id"]]
        distribution = {key: sum(1 for answer in current if answer.get("selected_option") == key) for key in "ABCD"}
        correct_count = sum(1 for answer in current if answer.get("is_correct") is True)
        answered_count = sum(1 for answer in current if answer.get("selected_option") or answer.get("is_correct") is not None)
        text_answers = []
        if question["question_type"] == "text_input":
            for answer in current:
                if not (answer.get("text_answer") or "").strip():
                    continue
                profile = profile_map.get(answer.get("student_id"), {})
                text_answers.append({
                    "student_id": answer.get("student_id"), "name": profile.get("name"),
                    "student_number": profile.get("student_number"), "class": profile.get("classes"),
                    "answer": answer.get("text_answer"), "submitted_at": answer.get("answered_at"),
                })
        reports.append({
            "id": question["id"], "sort_order": question["sort_order"], "question_type": question["question_type"],
            "correct_option": question.get("correct_option"), "submitted_count": answered_count,
            "unanswered_count": max(len(participant_ids) - answered_count, 0), "correct_count": correct_count,
            "accuracy": round(correct_count / answered_count * 100, 1) if answered_count and question["question_type"] == "single_choice" else 0,
            "distribution": distribution, "text_answers": text_answers,
        })
    topic = batches[0].get("topics") or {}
    return {
        "title": f"{topic.get('code', '')} {topic.get('name', '')}",
        "class": classes[0], "batch_id": batch_id, "student_count": len(students),
        "completed_count": len(participant_ids), "participant_count": len(participant_ids), "questions": reports,
    }


@app.get("/api/teacher/reports/{kind}/{item_id}")
async def teacher_report(kind: str, item_id: str, request: Request, authorization: str | None = Header(default=None)) -> dict[str, Any]:
    token = bearer_token(authorization); teacher = await current_teacher_profile(request, token)
    if kind == "classroom":
        status, sessions = await supabase_request(request, "GET", "/rest/v1/classroom_sessions", service_role=True, params={"id": f"eq.{item_id}", "created_by": f"eq.{teacher['user_id']}", "select": "id,batch_id,status,created_at,closed_at,class_id,classes(name,code)"})
        if status >= 400 or not sessions: raise HTTPException(status_code=404, detail="找不到课堂")
        session = sessions[0]; _, links = await supabase_request(request, "GET", "/rest/v1/classroom_session_questions", service_role=True, params={"session_id": f"eq.{item_id}", "select": "question_id,sort_order", "order": "sort_order.asc"})
        _, answers = await supabase_request(request, "GET", "/rest/v1/classroom_answers", service_role=True, params={"session_id": f"eq.{item_id}", "select": "question_id,selected_option,is_correct,student_id"})
        participant_status, participants = await supabase_request(request, "GET", "/rest/v1/classroom_participants", service_role=True, params={"session_id": f"eq.{item_id}", "select": "student_id"})
        title = "课堂统计"; total = len(participants) if participant_status < 400 else 0
    elif kind == "practice":
        status, practices = await supabase_request(request, "GET", "/rest/v1/practices", service_role=True, params={"id": f"eq.{item_id}", "created_by": f"eq.{teacher['user_id']}", "select": "id,batch_id,access_code,status,created_at"})
        if status >= 400 or not practices: raise HTTPException(status_code=404, detail="找不到练习")
        session = practices[0]
        _, links = await supabase_request(request, "GET", "/rest/v1/practice_questions", service_role=True, params={"practice_id": f"eq.{item_id}", "select": "question_id,sort_order", "order": "sort_order.asc"})
        round_status, rounds = await supabase_request(request, "GET", "/rest/v1/practice_rounds", service_role=True, params={"practice_id": f"eq.{item_id}", "status": "eq.completed", "select": "id,student_id,completed_at", "order": "completed_at.desc"})
        if round_status >= 400:
            raise HTTPException(status_code=502, detail="无法读取课后练习轮次")
        latest_round_ids: list[str] = []
        latest_students: set[str] = set()
        for round_row in rounds:
            student_id = round_row.get("student_id")
            if student_id and student_id not in latest_students:
                latest_students.add(student_id)
                latest_round_ids.append(round_row["id"])
        _, answers = await supabase_request(request, "GET", "/rest/v1/practice_answers", service_role=True, params={"round_id": f"in.({','.join(latest_round_ids)})", "state": "in.(confirmed,unanswered)", "select": "question_id,selected_option,is_correct,round_id"}) if latest_round_ids else (200, [])
        total = len(latest_round_ids)
        title = f"练习码 {session['access_code']}"
    else: raise HTTPException(status_code=404, detail="统计类型无效")
    question_ids = [link["question_id"] for link in links]
    _, questions = await supabase_request(request, "GET", "/rest/v1/questions", service_role=True, params={"id": f"in.({','.join(question_ids)})" if question_ids else "eq.invalid", "select": "id,sort_order,question_type,correct_option"})
    reports = []
    for question in sorted(questions, key=lambda row: row["sort_order"]):
        current = [answer for answer in answers if answer.get("question_id") == question["id"]]
        distribution = {key: sum(1 for answer in current if answer.get("selected_option") == key) for key in "ABCD"}
        correct = sum(1 for answer in current if answer.get("is_correct") is True)
        reports.append({"id": question["id"], "sort_order": question["sort_order"], "question_type": question["question_type"], "correct_option": question.get("correct_option"), "submitted_count": len(current), "correct_count": correct, "accuracy": round(correct / len(current) * 100) if current else 0, "distribution": distribution})
    return {"title": title, "session": session, "participant_count": total, "questions": reports}


@app.get("/api/student/words")
async def student_words(request: Request, grade_id: str, authorization: str | None = Header(default=None)) -> list[dict[str, Any]]:
    student = await verify_student_token(request, authorization)
    grade_status, grades = await supabase_request(request, "GET", "/rest/v1/grades", service_role=True, params={"code": f"eq.{grade_id.upper()}", "school_id": f"eq.{student['school_id']}", "is_active": "eq.true", "select": "id"})
    if grade_status >= 400 or not grades:
        raise HTTPException(status_code=404, detail="找不到可用年级")
    target_grade_id = grades[0]["id"]
    status, rows = await supabase_request(request, "GET", "/rest/v1/words", service_role=True, params={"grade_id": f"eq.{target_grade_id}", "select": "id,word,meaning,sort_order,word_batches!inner(status)", "word_batches.status": "eq.published", "order": "sort_order.asc"})
    if status >= 400: raise HTTPException(status_code=502, detail="无法读取单词")
    return rows


@app.get("/api/student/words/review")
async def student_words_review(request: Request, grade_id: str, authorization: str | None = Header(default=None)) -> list[dict[str, Any]]:
    student = await verify_student_token(request, authorization)
    grade_status, grades = await supabase_request(request, "GET", "/rest/v1/grades", service_role=True, params={"code": f"eq.{grade_id.upper()}", "school_id": f"eq.{student['school_id']}", "is_active": "eq.true", "select": "id"})
    if grade_status >= 400 or not grades:
        raise HTTPException(status_code=404, detail="找不到可用年级")
    target_grade_id = grades[0]["id"]
    status, rows = await supabase_request(request, "GET", "/rest/v1/word_progress", service_role=True, params={"student_id": f"eq.{student['id']}", "due_at": f"lte.{utc_now().isoformat()}", "words.grade_id": f"eq.{target_grade_id}", "select": "word_id,last_rating,interval_days,due_at,words!inner(id,word,meaning,grade_id)", "order": "due_at.asc"})
    if status >= 400: raise HTTPException(status_code=502, detail="无法读取复习单词")
    return rows


@app.post("/api/student/words/rate")
async def rate_student_word(payload: WordRatingRequest, request: Request, authorization: str | None = Header(default=None)) -> dict[str, Any]:
    student = await verify_student_token(request, authorization)
    status, result = await supabase_request(request, "POST", "/rest/v1/rpc/app_rate_word", service_role=True, body={"p_student": student["id"], "p_word": payload.word_id, "p_rating": payload.rating})
    if status >= 400: raise HTTPException(status_code=409, detail="单词进度保存失败")
    return result


@app.get("/api/teacher/words")
async def teacher_words(request: Request, authorization: str | None = Header(default=None)) -> list[dict[str, Any]]:
    token = bearer_token(authorization); teacher = await current_teacher_profile(request, token)
    status, rows = await supabase_request(request, "GET", "/rest/v1/word_batches", service_role=True, params={"status": "eq.published", "select": "id,grade_id,created_at,grades(code,name),words(count)", "grades.school_id": f"eq.{teacher['school_id']}"})
    if status >= 400: raise HTTPException(status_code=502, detail="无法读取单词词库")
    return [{**row, "word_count": (row.get("words") or [{"count": 0}])[0].get("count", 0)} for row in rows]


@app.post("/api/teacher/words/import")
async def import_words(request: Request, grade_id: str = Form(...), file: UploadFile = File(...), authorization: str | None = Header(default=None)) -> dict[str, Any]:
    token = bearer_token(authorization); teacher = await current_teacher_profile(request, token)
    if not file.filename or not file.filename.lower().endswith(".xlsx"): raise HTTPException(status_code=422, detail="只支持 .xlsx 格式的 Excel 文件")
    from openpyxl import load_workbook
    from io import BytesIO
    try: rows = list(load_workbook(filename=BytesIO(await file.read()), read_only=True, data_only=True).active.iter_rows(values_only=True))
    except Exception as exc: raise HTTPException(status_code=422, detail="无法读取 Excel 文件") from exc
    if not rows or len(rows[0]) < 2: raise HTTPException(status_code=422, detail="单词表必须包含单词、词义两列")
    headers = [str(value or '').strip().lower() for value in rows[0]]
    try: word_col = next(i for i, value in enumerate(headers) if value in {'单词','word'}) ; meaning_col = next(i for i, value in enumerate(headers) if value in {'词义','词意','meaning','中文'})
    except StopIteration as exc: raise HTTPException(status_code=422, detail="单词表必须包含“单词”和“词义”列") from exc
    words = []; seen = set(); errors = []
    for row_number, row in enumerate(rows[1:], start=2):
        word = str(row[word_col] or '').strip() if word_col < len(row) else ''; meaning = str(row[meaning_col] or '').strip() if meaning_col < len(row) else ''
        if not word and not meaning: continue
        normalized = word.casefold()
        if not word or not meaning: errors.append(f"第 {row_number} 行：单词和词义不能为空"); continue
        if normalized in seen: continue
        seen.add(normalized); words.append((word, meaning))
    if errors: raise HTTPException(status_code=422, detail={"message": "Excel 文件存在错误", "errors": errors})
    if not words: raise HTTPException(status_code=422, detail="没有可导入的单词")
    _, grades = await supabase_request(request, "GET", "/rest/v1/grades", service_role=True, params={"id": f"eq.{grade_id}", "school_id": f"eq.{teacher['school_id']}", "select": "id"})
    if not grades: raise HTTPException(status_code=404, detail="找不到年级")
    job_id, batch_id = str(uuid4()), str(uuid4())
    await supabase_request(request, "POST", "/rest/v1/import_jobs", service_role=True, body={"id": job_id, "created_by": teacher["user_id"], "kind": "words", "grade_id": grade_id, "filename": file.filename, "status": "preview"})
    batch_status, _ = await supabase_request(request, "POST", "/rest/v1/word_batches", service_role=True, body={"id": batch_id, "grade_id": grade_id, "created_by": teacher["user_id"], "import_job_id": job_id, "status": "draft"})
    if batch_status >= 400: raise HTTPException(status_code=502, detail="无法建立单词批次")
    word_rows = [{"batch_id": batch_id, "grade_id": grade_id, "word": word, "meaning": meaning, "sort_order": index} for index, (word, meaning) in enumerate(words, start=1)]
    write_status, _ = await supabase_request(request, "POST", "/rest/v1/words", service_role=True, body=word_rows)
    if write_status >= 400: raise HTTPException(status_code=502, detail="单词写入失败")
    publish_status, _ = await supabase_request(request, "POST", "/rest/v1/rpc/app_publish_word_batch", service_role=True, body={"p_batch": batch_id, "p_teacher": teacher["user_id"]})
    if publish_status >= 400: raise HTTPException(status_code=422, detail="单词发布失败")
    return {"batch_id": batch_id, "word_count": len(words)}


@app.get("/api/health")
async def health() -> dict[str, str]:
    return {"status": "ok", "service": "hhx-backend"}


@app.post("/api/auth/login")
async def login(payload: LoginRequest, request: Request) -> dict[str, Any]:
    username = payload.username.strip()
    if payload.mode == "student":
        status, students = await supabase_request(
            request,
            "POST",
            "/rest/v1/rpc/app_authenticate_student",
            service_role=True,
            body={"p_student_number": username, "p_password": payload.password},
        )
        if status >= 400 or not students:
            raise HTTPException(status_code=401, detail="學號或密碼不正確")
        student = students[0]
        student_token, expires_at = await create_student_session(request, student)
        return {
            "mode": "student",
            "student_token": student_token,
            "expires_at": expires_at.isoformat(),
            "user": {
                "id": student["student_id"],
                "student_number": student["student_number"],
                "display_name": student["name"],
                "grade_id": student["grade_id"],
                "grade": student["grade_code"],
                "class_id": student["class_id"],
                "class_name": student["class_name"],
            },
        }

    configured_email = env_value(request, "TEACHER_LOGIN_EMAIL", "admin@hhx.local")
    email = configured_email if username.lower() == "admin" else username
    if "@" not in email:
        raise HTTPException(status_code=401, detail="帳號或密碼不正確")

    status, result = await supabase_request(
        request,
        "POST",
        "/auth/v1/token",
        params={"grant_type": "password"},
        body={"email": email, "password": payload.password},
    )
    if status >= 400:
        raise HTTPException(status_code=401, detail="帳號或密碼不正確")
    return {
        "mode": "teacher",
        "access_token": result.get("access_token"),
        "expires_in": result.get("expires_in"),
        "user": result.get("user"),
    }


@app.get("/api/student/me")
async def current_student(
    request: Request,
    authorization: str | None = Header(default=None),
) -> dict[str, Any]:
    student = await verify_student_token(request, authorization)
    grade = await _student_grade(request, student)
    class_status, classes = await supabase_request(
        request, "GET", "/rest/v1/classes", service_role=True,
        params={"id": f"eq.{student['class_id']}", "school_id": f"eq.{student['school_id']}", "select": "id,name,code"},
    )
    if class_status >= 400 or not classes:
        raise HTTPException(status_code=403, detail="學生班級資料不可用")
    return {**student, "display_name": student["name"], "grade": grade["code"], "class_name": classes[0]["name"]}


async def _student_grade(request: Request, student: dict[str, Any]) -> dict[str, Any]:
    status, grades = await supabase_request(
        request,
        "GET",
        "/rest/v1/grades",
        service_role=True,
        params={"id": f"eq.{student['grade_id']}", "school_id": f"eq.{student['school_id']}", "select": "id,code,name,is_active"},
    )
    if status >= 400 or not grades or not grades[0].get("is_active"):
        raise HTTPException(status_code=403, detail="學生年級資料不可用")
    return grades[0]


async def _published_topic(request: Request, student: dict[str, Any], topic_id: str) -> tuple[dict[str, Any], dict[str, Any]]:
    status, topics = await supabase_request(
        request,
        "GET",
        "/rest/v1/topics",
        service_role=True,
        params={"id": f"eq.{topic_id}", "school_id": f"eq.{student['school_id']}", "is_active": "eq.true", "select": "id,school_id,grade_id,code,name"},
    )
    if status >= 400 or not topics:
        raise HTTPException(status_code=404, detail="找不到可用主題")
    grade = await _student_grade(request, student)
    if grade["code"] != "S6" or topics[0]["grade_id"] != student["grade_id"]:
        raise HTTPException(status_code=403, detail="自主數學練習僅開放 S6 學生")
    batch_status, batches = await supabase_request(
        request,
        "GET",
        "/rest/v1/math_batches",
        service_role=True,
        params={"topic_id": f"eq.{topic_id}", "status": "eq.published", "select": "id,topic_id,status", "limit": "1"},
    )
    if batch_status >= 400:
        raise HTTPException(status_code=502, detail="無法讀取主題內容")
    if not batches:
        raise HTTPException(status_code=404, detail="這個主題目前沒有題目")
    return topics[0], batches[0]


async def _question_payloads(
    request: Request,
    batch_id: str,
    answer_rows: list[dict[str, Any]] | None = None,
    *,
    practice_id: str | None = None,
) -> list[dict[str, Any]]:
    if practice_id:
        link_status, links = await supabase_request(
            request, "GET", "/rest/v1/practice_questions", service_role=True,
            params={"practice_id": f"eq.{practice_id}", "select": "question_id,sort_order", "order": "sort_order.asc"},
        )
        if link_status >= 400:
            raise HTTPException(status_code=502, detail="無法讀取練習題目清單")
        question_ids = [item["question_id"] for item in links]
        if not question_ids:
            return []
        query = {"id": f"in.({','.join(question_ids)})", "batch_id": f"eq.{batch_id}",
                 "select": "id,sort_order,question_type,image_bucket,image_path,image_bytes,correct_option", "order": "sort_order.asc"}
    else:
        query = {"batch_id": f"eq.{batch_id}",
                 "select": "id,sort_order,question_type,image_bucket,image_path,image_bytes,correct_option", "order": "sort_order.asc"}
    question_status, questions = await supabase_request(request, "GET", "/rest/v1/questions", service_role=True, params=query)
    if question_status >= 400:
        raise HTTPException(status_code=502, detail="無法讀取題目")
    answers = {item["question_id"]: item for item in (answer_rows or [])}
    payloads = []
    for question in questions:
        answer = answers.get(question["id"])
        image_url = await _signed_question_image(request, question["image_bucket"], question["image_path"])
        result = {
            "id": question["id"], "sort_order": question["sort_order"],
            "question_type": question["question_type"], "image_url": image_url,
            "answer": ({"selected_option": answer.get("selected_option"),
                        "text_answer": answer.get("text_answer"), "state": answer["state"],
                        "is_correct": answer.get("is_correct")} if answer else None),
        }
        if answer and answer.get("state") == "confirmed":
            result["correct_option"] = question.get("correct_option")
        payloads.append(result)
    return payloads


async def _signed_question_image(request: Request, bucket: str, path: str) -> str:
    from urllib.parse import quote
    status, result = await supabase_request(
        request, "POST", f"/storage/v1/object/sign/{quote(bucket, safe='')}/{quote(path, safe='/')}",
        service_role=True, body={"expiresIn": 900},
    )
    signed_url = result.get("signedURL") or result.get("signedUrl") if isinstance(result, dict) else None
    if status >= 400 or not signed_url:
        raise HTTPException(status_code=502, detail="無法讀取題目圖片")
    if signed_url.startswith("http://") or signed_url.startswith("https://"):
        return signed_url
    return f"{supabase_url(request)}{signed_url if signed_url.startswith('/storage/v1/') else '/storage/v1' + signed_url}"


@app.get("/api/student/topics")
async def list_student_topics(request: Request, authorization: str | None = Header(default=None)) -> list[dict[str, Any]]:
    student = await verify_student_token(request, authorization)
    grade = await _student_grade(request, student)
    if grade["code"] != "S6":
        raise HTTPException(status_code=403, detail="自主數學練習僅開放 S6 學生；請使用老師提供的練習碼")
    status, topics = await supabase_request(
        request, "GET", "/rest/v1/topics", service_role=True,
        params={"school_id": f"eq.{student['school_id']}", "grade_id": f"eq.{student['grade_id']}",
                "is_active": "eq.true", "select": "id,code,name,sort_order", "order": "sort_order.asc"},
    )
    if status >= 400:
        raise HTTPException(status_code=502, detail="無法讀取數學主題")
    if not topics:
        return []
    ids = [item["id"] for item in topics]
    batch_status, batches = await supabase_request(
        request, "GET", "/rest/v1/math_batches", service_role=True,
        params={"topic_id": f"in.({','.join(ids)})", "status": "eq.published", "select": "id,topic_id,questions(count)"},
    )
    if batch_status >= 400:
        raise HTTPException(status_code=502, detail="無法讀取主題題目數量")
    by_topic = {}
    for batch in batches:
        counts = batch.get("questions") or [{"count": 0}]
        by_topic[batch["topic_id"]] = {"batch_id": batch["id"], "question_count": counts[0].get("count", 0)}
    return [{**topic, **by_topic.get(topic["id"], {"batch_id": None, "question_count": 0})} for topic in topics]


@app.get("/api/student/topics/{topic_id}")
async def get_student_topic(topic_id: str, request: Request, authorization: str | None = Header(default=None)) -> dict[str, Any]:
    student = await verify_student_token(request, authorization)
    topic, batch = await _published_topic(request, student, topic_id)
    questions = await _question_payloads(request, batch["id"])
    return {**topic, "batch_id": batch["id"], "questions": questions}


async def _student_practice_context(request: Request, student: dict[str, Any], access_code: str) -> tuple[dict[str, Any], dict[str, Any]]:
    status, practices = await supabase_request(
        request, "GET", "/rest/v1/practices", service_role=True,
        params={"access_code": f"eq.{access_code.strip().upper()}", "status": "eq.active", "select": "id,batch_id,access_code,status", "limit": "1"},
    )
    if status >= 400 or not practices:
        raise HTTPException(status_code=404, detail="練習碼無效或已失效")
    practice = practices[0]
    batch_status, batches = await supabase_request(
        request, "GET", "/rest/v1/math_batches", service_role=True,
        params={"id": f"eq.{practice['batch_id']}", "status": "eq.published", "select": "id,topic_id,status,topics(id,school_id,grade_id,code,name,is_active)"},
    )
    if batch_status >= 400 or not batches:
        raise HTTPException(status_code=404, detail="練習內容已更新或不可用")
    topic = batches[0].get("topics")
    if not topic or not topic.get("is_active") or topic["school_id"] != student["school_id"]:
        raise HTTPException(status_code=403, detail="此練習碼不可用")
    return practice, batches[0]


async def _start_student_round(request: Request, student: dict[str, Any], batch: dict[str, Any], practice_id: str | None) -> dict[str, Any]:
    params = {"student_id": f"eq.{student['id']}", "batch_id": f"eq.{batch['id']}",
              "status": "eq.in_progress", "select": "id,student_id,batch_id,practice_id,status"}
    params["practice_id"] = f"eq.{practice_id}" if practice_id else "is.null"
    status, rounds = await supabase_request(request, "GET", "/rest/v1/practice_rounds", service_role=True, params=params)
    if status >= 400:
        raise HTTPException(status_code=502, detail="無法恢復練習進度")
    if rounds:
        return rounds[0]
    status, created = await supabase_request(
        request, "POST", "/rest/v1/practice_rounds", service_role=True, prefer_representation=True,
        body={"student_id": student["id"], "batch_id": batch["id"], "practice_id": practice_id, "status": "in_progress"},
    )
    if status >= 400 or not created:
        raise HTTPException(status_code=409, detail="無法開始練習；內容可能已更新")
    return created[0]


async def _round_payload(request: Request, round_row: dict[str, Any], practice_id: str | None) -> dict[str, Any]:
    answer_status, answers = await supabase_request(
        request, "GET", "/rest/v1/practice_answers", service_role=True,
        params={"round_id": f"eq.{round_row['id']}", "select": "question_id,selected_option,text_answer,state,is_correct,confirmed_at", "order": "updated_at.asc"},
    )
    if answer_status >= 400:
        raise HTTPException(status_code=502, detail="無法恢復答案")
    questions = await _question_payloads(request, round_row["batch_id"], answers, practice_id=practice_id)
    return {"round_id": round_row["id"], "status": round_row["status"], "questions": questions}


@app.post("/api/student/rounds/start")
async def start_student_round(payload: StudentRoundStartRequest, request: Request,
                              authorization: str | None = Header(default=None)) -> dict[str, Any]:
    student = await verify_student_token(request, authorization)
    if bool(payload.topic_id) == bool(payload.access_code):
        raise HTTPException(status_code=422, detail="必須指定主題或練習碼其中一種")
    if payload.topic_id:
        _, batch = await _published_topic(request, student, payload.topic_id)
        practice_id = None
    else:
        practice, batch = await _student_practice_context(request, student, payload.access_code or "")
        practice_id = practice["id"]
    round_row = await _start_student_round(request, student, batch, practice_id)
    return await _round_payload(request, round_row, practice_id)


@app.post("/api/student/rounds/answer")
async def save_student_round_answer(payload: StudentRoundAnswerRequest, request: Request,
                                    authorization: str | None = Header(default=None)) -> dict[str, Any]:
    student = await verify_student_token(request, authorization)
    status, result = await supabase_request(
        request, "POST", "/rest/v1/rpc/app_save_practice_answer", service_role=True,
        body={"p_round": payload.round_id, "p_student": student["id"], "p_question": payload.question_id,
              "p_option": payload.selected_option, "p_text": payload.text_answer, "p_confirm": payload.confirm},
    )
    if status >= 400:
        raise HTTPException(status_code=409, detail="答案保存失敗；練習可能已結束或題目已更新")
    return result if isinstance(result, dict) else {"saved": True}


@app.post("/api/student/rounds/finish")
async def finish_student_round(payload: StudentRoundFinishRequest, request: Request,
                               authorization: str | None = Header(default=None)) -> Any:
    student = await verify_student_token(request, authorization)
    status, result = await supabase_request(
        request, "POST", "/rest/v1/rpc/app_finish_round", service_role=True,
        body={"p_round": payload.round_id, "p_student": student["id"],
              "p_answers": [item.model_dump() for item in payload.answers]},
    )
    if status >= 400:
        raise HTTPException(status_code=409, detail="整份提交失敗；練習可能已結束或題目已更新")
    return {"completed": True, "answers": result}


@app.post("/api/student/logout")
async def logout_student(
    request: Request,
    authorization: str | None = Header(default=None),
) -> dict[str, bool]:
    token = bearer_token(authorization)
    status, _ = await supabase_request(
        request,
        "PATCH",
        "/rest/v1/student_sessions",
        service_role=True,
        params={
            "token_hash": f"eq.{hashlib.sha256(token.encode('utf-8')).hexdigest()}",
            "revoked_at": "is.null",
        },
        body={"revoked_at": utc_now().isoformat()},
    )
    if status >= 400:
        raise HTTPException(status_code=502, detail="學生登出失敗")
    return {"logged_out": True}


@app.get("/api/auth/me")
async def current_teacher(
    request: Request,
    authorization: str | None = Header(default=None),
) -> dict[str, Any]:
    token = bearer_token(authorization)
    profile = await current_teacher_profile(request, token)
    return profile



Default = asgi.entrypoint(app)
