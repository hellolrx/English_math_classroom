"""Vocabulary Excel parser for importing English words."""

from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from typing import Any

from openpyxl import load_workbook


MAX_WORD_COUNT = 5000


@dataclass(frozen=True)
class ParsedWord:
    row_number: int
    word: str
    meaning: str


class VocabularyImportError(ValueError):
    def __init__(self, errors: list[str]):
        self.errors = errors
        super().__init__("Excel 文件存在错误，请修正后重新上传")


def _text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _normalise_header(value: Any) -> str:
    return "".join(_text(value).lower().split())


VOCAB_HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "word": ("单词", "單詞", "word", "words", "english"),
    "meaning": ("词义", "詞義", "意思", "meaning", "translation", "chinese"),
}


def _find_columns(header_row: tuple[Any, ...]) -> dict[str, int]:
    normalized = {_normalise_header(value): index for index, value in enumerate(header_row)}
    columns: dict[str, int] = {}
    missing: list[str] = []
    for field, aliases in VOCAB_HEADER_ALIASES.items():
        for alias in aliases:
            index = normalized.get(_normalise_header(alias))
            if index is not None:
                columns[field] = index
                break
        if field not in columns:
            missing.append(aliases[0])
    if missing:
        raise VocabularyImportError([f"缺少必要欄位：{', '.join(missing)}"])
    return columns


def parse_vocabulary_excel(content: bytes) -> list[ParsedWord]:
    """Parse vocabulary Excel file with '单词' and '词义' columns."""
    if not content:
        raise VocabularyImportError(["Excel 文件为空"])
    
    try:
        workbook = load_workbook(filename=BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise VocabularyImportError(["无法读取 Excel 文件，请确认文件格式为 .xlsx"]) from exc
    
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    
    if not rows:
        raise VocabularyImportError(["Excel 文件没有内容"])
    
    columns = _find_columns(tuple(rows[0]))
    parsed: list[ParsedWord] = []
    errors: list[str] = []
    seen_words: set[str] = set()
    
    for row_number, row in enumerate(rows[1:], start=2):
        values = tuple(row)
        if not any(_text(value) for value in values):
            continue
        
        if len(parsed) >= MAX_WORD_COUNT:
            errors.append(f"第 {row_number} 行之后单词超过上限 {MAX_WORD_COUNT} 个")
            break
        
        word = _text(values[columns["word"]] if columns["word"] < len(values) else None)
        meaning = _text(values[columns["meaning"]] if columns["meaning"] < len(values) else None)
        
        row_errors: list[str] = []
        if not word:
            row_errors.append("单词不能为空")
        if not meaning:
            row_errors.append("词义不能为空")
        
        if word and word.lower() in seen_words:
            row_errors.append(f"单词 '{word}' 重复")
        
        if row_errors:
            errors.append(f"第 {row_number} 行：{'；'.join(row_errors)}")
            continue
        
        seen_words.add(word.lower())
        parsed.append(ParsedWord(row_number, word, meaning))
    
    if not parsed and not errors:
        errors.append("Excel 文件没有可导入的单词")
    if errors:
        raise VocabularyImportError(errors)
    
    return parsed
